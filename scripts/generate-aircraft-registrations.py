#!/usr/bin/env python3
"""Generate aircraft lookup data from the combined British Airways fleet workbook.

Requires openpyxl for this maintenance script only. The workbook is the source
for registration details and aircraft summaries; AOR position data remains in
the individual aircraft JavaScript files.
"""

from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


SOURCE_NAME = "British Airways Fleet.xlsx"
SOURCE_AS_OF = "2026-09"
SHEET_NAME = "Registrations"
EXCLUDED_AIRFILES = {"78Z"}
LONGHAUL_AIRFILES = {"38A", "38T", "351", "78E", "78N", "789", "781", "77M", "77L", "77H", "77T", "77S"}
AIRFILE_MARKER_START = "// BEGIN GENERATED AIRFILE DATA"
AIRFILE_MARKER_END = "// END GENERATED AIRFILE DATA"
EQUIPMENT_COLUMNS = (
    ("AED Location", "AED"),
    ("M5 Location", "M5"),
    ("RESUS Location", "RES"),
    ("RESTRAINT Location", "RESTRAINT"),
    ("FE Location", "FE"),
    ("WEX Location", "WEX"),
)


def normalise_header(value: object) -> str:
    return " ".join(str(value or "").strip().casefold().replace("–", "-").replace("—", "-").split())


def value_for(row: tuple, headers: dict[str, int], *names: str):
    for name in names:
        index = headers.get(normalise_header(name))
        if index is not None and index < len(row):
            return row[index]
    return None


def clean(value):
    """Treat workbook NA, empty strings and whitespace as not applicable."""
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        if not value or value.casefold() == "na":
            return None
    return value


def yes_no(value: object) -> bool | None:
    value = clean(value)
    if value is None:
        return None
    normalised = str(value).casefold()
    if normalised in {"yes", "y", "true", "1"}:
        return True
    if normalised in {"no", "n", "false", "0"}:
        return False
    raise ValueError(f"Unexpected yes/no value: {value!r}")


def equipment_for(row: tuple, headers: dict[str, int]) -> list[dict]:
    equipment = []
    for column, code in EQUIPMENT_COLUMNS:
        location = clean(value_for(row, headers, column))
        if location is not None:
            equipment.append({"code": code, "location": str(location)})
    return equipment


def spare_seats(count, raw_location) -> list[dict]:
    count, raw_location = clean(count), clean(raw_location)
    if not count or not raw_location:
        return []
    locations = [part.strip() for part in re.split(r"\s*,\s*|\s*&\s*", str(raw_location)) if part.strip()]
    results = []
    for location in locations[:int(count)]:
        match = re.search(r"\b([A-Z]?\d+[A-Z]?)\b", location.upper())
        if not match:
            results.append({"seat": location, "facing": None, "note": "Spare crew seat"})
            continue
        seat = match.group(1)
        facing_match = re.search(r"\b(FWD|AFT)\b", location.upper())
        descriptor_match = re.search(r"\b(INBOARD|CENTRE|CENTER)\b", location.upper())
        item = {"seat": seat, "facing": facing_match.group(1) if facing_match else None, "note": "Spare crew seat"}
        if descriptor_match:
            item["descriptor"] = descriptor_match.group(1).title()
        results.append(item)
    return results


def cabin_breakdown(row: tuple, headers: dict[str, int]) -> tuple[dict, list[str]]:
    breakdown = {}
    classes = []
    for header, key, code in (
        ("F - First", "first", "F"),
        ("J - Club World", "clubWorld", "J"),
        ("W - World Traveller Plus", "worldTravellerPlus", "W"),
        ("M - World Traveller", "worldTraveller", "M"),
    ):
        count = clean(value_for(row, headers, header))
        if count is None:
            continue
        if float(count) > 0:
            breakdown[key] = count
            classes.append(code)
    return breakdown, classes


def read_workbook(workbook_path: Path) -> tuple[dict[str, dict], dict[str, list[dict]]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"Workbook must contain a '{SHEET_NAME}' sheet")
    rows = workbook[SHEET_NAME].iter_rows(values_only=True)
    next(rows, None)  # workbook title row
    header_row = next(rows, None)
    if not header_row:
        raise ValueError(f"No header row found on the '{SHEET_NAME}' sheet")
    headers = {normalise_header(value): index for index, value in enumerate(header_row) if value is not None}

    raw_rows: dict[str, list[tuple]] = defaultdict(list)
    registrations: dict[str, dict] = {}
    for row in rows:
        registration = clean(value_for(row, headers, "Registration"))
        airfile_value = clean(value_for(row, headers, "Airfile"))
        if registration is None or airfile_value is None:
            continue
        registration = str(registration).upper()
        airfile = str(airfile_value).strip()
        if airfile in EXCLUDED_AIRFILES:
            continue
        if registration in registrations:
            raise ValueError(f"Duplicate registration found in workbook: {registration}")
        haul = "longhaul" if airfile in LONGHAUL_AIRFILES else "shorthaul"
        record = {
            "registration": registration,
            "haul": haul,
            "baseSection": clean(value_for(row, headers, "Base / section")),
            "airfile": airfile,
            "wifiType": clean(value_for(row, headers, "Wi-Fi Type")),
            "livery": clean(value_for(row, headers, "Livery")),
            "newShorthaulSeat": yes_no(value_for(row, headers, "New SH seat")),
            "xlOverheadBins": yes_no(value_for(row, headers, "XL overhead bins")),
            "sepEquipment": equipment_for(row, headers),
        }
        if haul == "longhaul":
            record["firstProduct"] = clean(value_for(row, headers, "First Product"))
            record["clubWorldProduct"] = clean(value_for(row, headers, "Club World Product"))
        registrations[registration] = {key: value for key, value in record.items() if value is not None}
        raw_rows[airfile].append(row)

    summaries: dict[str, list[dict]] = {"longhaul": [], "shorthaul": []}
    for airfile, type_rows in raw_rows.items():
        haul = "longhaul" if airfile in LONGHAUL_AIRFILES else "shorthaul"
        row = type_rows[0]
        variant = str(clean(value_for(row, headers, "Aircraft / variant")) or airfile)
        family_match = re.search(r"[AB]\d{3}", variant.upper())
        family = family_match.group(0) if family_match else variant
        manufacturer = str(clean(value_for(row, headers, "Manufacturer")) or ("Airbus" if family.startswith("A") else "Boeing"))
        equipment_by_registration = [equipment_for(item, headers) for item in type_rows]
        equipment_signatures = {json.dumps(items, sort_keys=True) for items in equipment_by_registration}
        equipment_varies = len(equipment_signatures) > 1
        shared_equipment = equipment_by_registration[0] if not equipment_varies else []
        total_seats = clean(value_for(row, headers, "Total Passenger Seats"))

        if haul == "longhaul":
            breakdown, classes = cabin_breakdown(row, headers)
            legal_minimum = clean(value_for(row, headers, "Legal Minimum Crew"))
            flight_bunks = yes_no(value_for(row, headers, "FC Bunks"))
            cabin_bunks = yes_no(value_for(row, headers, "CC Bunks"))
            rest_types = []
            if flight_bunks:
                rest_types.append("OFCR" if manufacturer.casefold() == "airbus" else "FCRC")
            if cabin_bunks:
                rest_types.append("OFAR" if manufacturer.casefold() == "airbus" else "CCRC")
            rest_lookup = {"OFCR": "flightCrewRest", "FCRC": "flightCrewRest", "OFAR": "cabinCrewRest", "CCRC": "cabinCrewRest"}
            crew = None
            if legal_minimum is not None:
                crew = {
                    "legalMinimum": legal_minimum,
                    "requiredSeats": list(range(1, int(legal_minimum) + 1)),
                    "totalCrewSeats": clean(value_for(row, headers, "Crew Seats")),
                    "standardCrewCompliment": clean(value_for(row, headers, "Normal Crew Complement")),
                    "spareCrewSeats": spare_seats(value_for(row, headers, "Number of Spare Crew Seats"), value_for(row, headers, "Location of Spare Crew Seats")),
                }
                crew = {key: value for key, value in crew.items() if value is not None}
            summary = {
                "airfile": airfile, "variant": variant, "family": family, "manufacturer": manufacturer,
                "configName": clean(value_for(row, headers, "Configuration")) or f"{len(classes)} Class",
                "classCount": len(classes), "classes": classes, "seatBreakdown": breakdown,
                "totalSeats": total_seats, "crew": crew, "includesRestFacilities": bool(rest_types),
                "restTypes": rest_types,
                "flightCrewRest": next((code for code in rest_types if rest_lookup[code] == "flightCrewRest"), None),
                "cabinCrewRest": next((code for code in rest_types if rest_lookup[code] == "cabinCrewRest"), None),
                "emergencyEquipmentSummary": shared_equipment,
                "registrationRequiredForEquipment": equipment_varies,
            }
        else:
            cabin_values = str(clean(value_for(row, headers, "Cabin Codes")) or "J, M").replace(" ", "").split(",")
            classes = [code for code in cabin_values if code in {"F", "J", "W", "M"}]
            summary = {
                "airfile": airfile, "variant": variant, "family": family, "manufacturer": manufacturer,
                "configName": clean(value_for(row, headers, "Configuration")) or "Shorthaul",
                "classCount": len(classes), "classes": classes, "seatBreakdown": None,
                "totalSeats": total_seats, "crew": None, "includesRestFacilities": False,
                "restTypes": [], "flightCrewRest": None, "cabinCrewRest": None,
                "emergencyEquipmentSummary": shared_equipment,
                "registrationRequiredForEquipment": equipment_varies,
            }
        summaries[haul].append(summary)

    workbook.close()
    return registrations, summaries


def update_aircraft_file(project_root: Path, summaries: dict[str, list[dict]]) -> None:
    aircraft_path = project_root / "data" / "aircraft.js"
    assignments = []
    longhaul_codes, shorthaul_codes = [], []
    for haul in ("longhaul", "shorthaul"):
        codes = longhaul_codes if haul == "longhaul" else shorthaul_codes
        for summary in summaries[haul]:
            code = summary["airfile"]
            codes.append(code)
            selector_sub = f'{summary["variant"]} · {summary["configName"]}' if haul == "longhaul" else f'{summary["variant"]} · {summary["totalSeats"]} seats'
            generated = {
                "code": "SH Airbus" if haul == "shorthaul" and summary["manufacturer"] == "Airbus" else code,
                "airfile": code, "haul": haul, "fullName": summary["variant"], "shortName": summary["family"],
                "family": summary["family"], "variant": summary["variant"], "manufacturer": summary["manufacturer"],
                "configName": summary["configName"], "classCount": summary["classCount"], "classes": summary["classes"],
                "includesRestFacilities": summary["includesRestFacilities"], "restTypes": summary["restTypes"],
                "flightCrewRest": summary["flightCrewRest"], "cabinCrewRest": summary["cabinCrewRest"],
                "totalSeats": summary["totalSeats"], "seatBreakdown": summary["seatBreakdown"],
                "crew": summary["crew"], "emergencyEquipmentSummary": summary["emergencyEquipmentSummary"],
                "registrationRequiredForEquipment": summary["registrationRequiredForEquipment"],
                # 38T has worksheet lookup data, but no AOR/FDP dataset yet.
                "dataFile": f"{code}.js" if haul == "longhaul" and code != "38T" else None,
                "dataPath": f"data/aircraft/{code}.js" if haul == "longhaul" and code != "38T" else None,
                "briefingTitle": f'{summary["variant"]} Briefing',
                "selectorLabel": code,
                "selectorSubLabel": selector_sub,
            }
            assignments.append(f"globalThis.AIRCRAFT[{json.dumps(code)}] = {json.dumps(generated, ensure_ascii=False, indent=2)};")
    # Keep 38T available for registration/type lookups without offering it as a
    # Briefing or FDP selection until a corresponding operational dataset exists.
    fdp_longhaul_codes = [code for code in longhaul_codes if code != "38T"]
    all_codes = longhaul_codes + shorthaul_codes
    block = "/* Generated from the fleet workbook by scripts/generate-aircraft-registrations.py. */\n"
    block += "globalThis.AIRCRAFT = {};\n\n" + AIRFILE_MARKER_START + "\n"
    block += "// Generated aircraft summaries share a single source: the fleet workbook.\n"
    block += "\n".join(assignments) + "\n"
    block += "globalThis.AIRCRAFT_ORDER_LONGHAUL = " + json.dumps(fdp_longhaul_codes) + ";\n"
    block += "globalThis.AIRCRAFT_ORDER_SHORTHAUL = " + json.dumps(shorthaul_codes) + ";\n"
    block += "globalThis.AIRCRAFT_ORDER = " + json.dumps(all_codes) + ";\n"
    block += "globalThis.AIRFILE_ORDER = " + json.dumps(all_codes) + ";\n"
    block += AIRFILE_MARKER_END + "\n"
    aircraft_path.write_text(block, encoding="utf-8", newline="\n")


def remove_shared_summary_from_aor(project_root: Path, summaries: dict[str, list[dict]]) -> None:
    """Keep workbook-owned aircraft summary data in aircraft.js, not duplicated in AOR files."""
    for summary in summaries["longhaul"]:
        path = project_root / "data" / "aircraft" / f'{summary["airfile"]}.js'
        if not path.exists():
            continue
        source = path.read_text(encoding="utf-8")
        source = re.sub(r'(?m)^\s{2}(?:"variantNotes"|variantNotes):.*?,\r?\n', "", source)
        source = re.sub(r'(?m)^\s{2}(?:"aircraftName"|aircraftName):.*?,\r?\n', "", source)
        source = re.sub(r'(?m)^\s{2}(?:"configName"|configName):.*?,\r?\n', "", source)
        source = re.sub(r'(?ms)^\s{2}(?:"config"|config):\s*\{.*?^\s{2}\},\r?\n', "", source)
        source = re.sub(r'(?ms)^\s{2}(?:"crew"|crew):\s*\{.*?^\s{2}\},\r?\n', "", source)
        source = re.sub(r'(?ms)^[ \t]{4}(?:"emergencyEquipmentSummary"|emergencyEquipmentSummary):[ \t]*\[.*?^[ \t]{4}\],[ \t]*\r?\n', "", source)
        path.write_text(source, encoding="utf-8", newline="\n")


def main() -> None:
    project_root = Path(__file__).resolve().parents[1]
    default_workbook = project_root.parent / SOURCE_NAME
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", nargs="?", type=Path, default=default_workbook)
    parser.add_argument("--output", type=Path, default=project_root / "data" / "aircraft-registrations.js")
    args = parser.parse_args()
    if not args.workbook.is_file():
        raise SystemExit(f"Workbook not found: {args.workbook}")

    records, summaries = read_workbook(args.workbook)
    update_aircraft_file(project_root, summaries)
    remove_shared_summary_from_aor(project_root, summaries)
    registration_order = list(records)
    output = f'''/* Generated from the BA fleet register. Run scripts/generate-aircraft-registrations.py to refresh. */
globalThis.AIRCRAFT_REGISTRATION_SOURCE = Object.freeze({{
  file: {json.dumps(SOURCE_NAME)},
  asOf: {json.dumps(SOURCE_AS_OF)},
  registrationCount: {len(records)}
}});

globalThis.AIRCRAFT_REGISTRATIONS = Object.freeze({json.dumps(records, ensure_ascii=False, indent=2)});
globalThis.AIRCRAFT_REGISTRATION_ORDER = Object.freeze({json.dumps(registration_order, ensure_ascii=False, indent=2)});

globalThis.getAircraftByRegistration = function (value) {{
  const compact = String(value || "").trim().toUpperCase().replace(/[^A-Z0-9]/g, "");
  const candidates = compact.startsWith("G") ? [compact, compact.slice(1)] : [compact];
  const registration = candidates.find((candidate) => globalThis.AIRCRAFT_REGISTRATIONS[candidate]);
  if (!registration) return null;
  const record = globalThis.AIRCRAFT_REGISTRATIONS[registration];
  return {{ ...record, aircraft: globalThis.AIRCRAFT?.[record.airfile] || null }};
}};
'''
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(output, encoding="utf-8", newline="\n")
    print(f"Wrote {len(records)} registrations and summaries for {sum(len(items) for items in summaries.values())} aircraft types")


if __name__ == "__main__":
    main()
